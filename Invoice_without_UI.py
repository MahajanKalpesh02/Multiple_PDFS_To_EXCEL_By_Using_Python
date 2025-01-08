import pdfplumber
import openpyxl
import re
import os

# Function to read existing data from the text file and check for duplicates
def check_duplicate_in_text_file(data, text_file_path):
    if not os.path.exists(text_file_path):
        return False  # File doesn't exist, so no duplicates
    with open(text_file_path, 'r') as file:
        existing_data = file.read()
        if data in existing_data:
            return True  # Data already exists
    return False

# Function to write the new data to the text file
def append_to_text_file(data, text_file_path):
    with open(text_file_path, 'a') as file:
        file.write(data + '\n')  # Append the new data to the file

# Open the PDF file
pdf_path = "your path .pdf"
with pdfplumber.open(pdf_path) as pdf:
    text = ''
    for page in pdf.pages:
        text += page.extract_text()

# Define the placeholders and patterns
columns = [
            
            # Give Your Placeholder/Column here......
            # eg. "Invoice Number"
            
]

details_patterns = {

    # Give your Pattern here.....
    #eg. "Invoice Number": r"Invoice(?: No)?[:\s]*([A-Za-z0-9\-/]+)"      # Pattern for Invoice Number
}

# Path for the Excel file and text file
excel_file_path = 'extracted_data.xlsx'
text_file_path = 'extracted_data.txt'

# Check if Excel file exists, and load or create workbook
if os.path.exists(excel_file_path):
    wb = openpyxl.load_workbook(excel_file_path)
    ws = wb.active
else:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(columns)  # Add headers if creating a new workbook

# Extract values and add to Excel
data = []
for column in columns:
    if column in details_patterns:
        match = re.search(details_patterns[column], text, re.IGNORECASE)
        if match:
            value = match.group(1)  # Capture the matched value
            data.append(value.strip())
        else:
            data.append("Not Found")
    else:
        # For columns not requiring regex, extract values using text parsing
        start_index = text.find(column)
        if start_index != -1:
            start_value = text[start_index + len(column):].strip()
            end_index = start_value.find("\n")
            value = start_value[:end_index] if end_index != -1 else start_value
            value = value.replace(":", "").strip()
            data.append(value)
        else:
            data.append("Not Found")

# Convert the extracted data to a string format to store in the text file
data_string = " | ".join(data)

# Check if this data is a duplicate
if not check_duplicate_in_text_file(data_string, text_file_path):
    # Append the new data to Excel
    ws.append(data)
    wb.save(excel_file_path)  # Save changes to the Excel file

    # Append to the text file
    append_to_text_file(data_string, text_file_path)
    print("New data added to Excel and text file.")
else:
    print("Data already exists in the text file. No duplicate added.")
